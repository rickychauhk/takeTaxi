#!/usr/bin/env python3
"""Generate UAT sign-off (簽收書) Excel + Word for Taxi MA."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent
TODAY = date.today().strftime("%Y-%m-%d")

# --- 交付版本（依 app.config / package.json 與已知建置；iOS 待上架時 build 欄留空供填寫）---
VERSION_ROWS: list[dict[str, str]] = [
    {
        "component": "後端 API",
        "platform": "Server (Docker)",
        "app_version": "0.1.0",
        "build_artifact": "taxima-api 映像（UAT/PROD 已部署）",
        "uat_url": "https://taxima.hkcodai.space/api/v1/health",
        "prod_url": "https://app.hktaketaxi.online/api/v1/health",
        "build_date": TODAY,
        "remark": "含貨主柯打大廳、發單者取消、WebSocket 等 UAT 變更",
    },
    {
        "component": "CRM / 後台",
        "platform": "Web (Vite)",
        "app_version": "0.1.0",
        "build_artifact": "taxima-crm 靜態站（容器重建）",
        "uat_url": "https://taxima.hkcodai.space",
        "prod_url": "https://app.hktaketaxi.online",
        "build_date": TODAY,
        "remark": "新增公開頁 /privacy、/support（App Store 審核用）",
    },
    {
        "component": "乘客 App",
        "platform": "Android (arm64-v8a)",
        "app_version": "1.0.0",
        "build_artifact": "dist/apk/taxima-passenger-prod-arm64.apk",
        "uat_url": "API: UAT/PROD 可切換（建置預設 PROD API）",
        "prod_url": "online.hktaketaxi.app",
        "build_date": "2026-10-29",
        "remark": "Debug 簽名內測包；Expo updates 關閉",
    },
    {
        "component": "司機 App",
        "platform": "Android (arm64-v8a)",
        "app_version": "1.0.0",
        "build_artifact": "dist/apk/taxima-driver-prod-arm64.apk",
        "uat_url": "API: UAT/PROD 可切換（建置預設 PROD API）",
        "prod_url": "online.hktaketaxi.driverapp",
        "build_date": "2026-10-29",
        "remark": "若含「取消我出嘅柯打／大廳顏色」等最新改動，請以最新 APK 建置日期為準",
    },
    {
        "component": "乘客 App",
        "platform": "iOS",
        "app_version": "1.0.0",
        "build_artifact": "（待 TestFlight / App Store Connect build 號）",
        "uat_url": "Bundle: online.hktaketaxi.app",
        "prod_url": "—",
        "build_date": "",
        "remark": "尚未提交商店建置；UAT 以實機 dev/Ad Hoc 為準（如有）",
    },
    {
        "component": "司機 App",
        "platform": "iOS",
        "app_version": "1.0.0",
        "build_artifact": "（待 TestFlight / App Store Connect build 號）",
        "uat_url": "Bundle: online.hktaketaxi.driverapp",
        "prod_url": "—",
        "build_date": "",
        "remark": "尚未提交商店建置；UAT 以實機 dev/Ad Hoc 為準（如有）",
    },
]

SCOPE_ITEMS: list[str] = [
    "CRM：營運管理、司機審批、叫車選項、積分與推播等既有功能；公開私隱政策與支援頁。",
    "API：叫車／接單／出貨全流程；貨主柯打於大廳對發單司機可見（「我出嘅」）；發單者不可自接；matching 狀態可取消自己出的柯打。",
    "司機 App：大廳列表（千足金／九折定食／公司價視覺區分）、出貨成功表單重置、貨主柯打留大廳。",
    "乘客 App：叫車、行程、訊息、失物協助等既有功能（本輪 UAT 以 Android 建置為主）。",
    "iOS 雙端：功能對齊 Android（以實際提交 TestFlight 之 build 為驗收標的）。",
]

UAT_ENV = {
    "UAT 主網域": "https://taxima.hkcodai.space",
    "UAT API Health": "https://taxima.hkcodai.space/api/v1/health",
    "UAT 私隱政策": "https://taxima.hkcodai.space/privacy",
    "UAT 支援頁": "https://taxima.hkcodai.space/support",
    "PROD 主網域": "https://app.hktaketaxi.online",
    "PROD 私隱政策": "https://app.hktaketaxi.online/privacy",
    "PROD 支援頁": "https://app.hktaketaxi.online/support",
}

VERSION_HEADERS = [
    "元件",
    "平台",
    "版本號",
    "建置產物／識別",
    "UAT／識別",
    "PROD／Bundle",
    "建置日期",
    "備註",
]

SIGNOFF_HEADERS = ["項目", "說明", "結果／填寫"]

THIN = Side(style="thin", color="999999")


def _header_style(cell, bold=True, fg="C00000", font_color="FFFFFF") -> None:
    cell.font = Font(bold=bold, color=font_color)
    cell.fill = PatternFill("solid", fgColor=fg)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def write_excel(path: Path) -> None:
    wb = Workbook()

    # Sheet 1: 簽收書摘要
    ws = wb.active
    ws.title = "UAT簽收書"
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "Taxi MA（搭的）用戶驗收測試（UAT）簽收書"
    c.font = Font(bold=True, size=16)
    c.alignment = Alignment(horizontal="center")

    ws["A3"] = "文件日期"
    ws["B3"] = TODAY
    ws["A4"] = "驗收環境"
    ws["B4"] = "UAT（主要）＋ 對照 PROD 網址"
    ws["A5"] = "Monorepo 標記版本"
    ws["B5"] = "0.1.0-uat"

    row = 7
    ws.cell(row=row, column=1, value="一、交付版本清單").font = Font(bold=True, size=12)
    row += 1
    for col, h in enumerate(VERSION_HEADERS, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        _header_style(cell)
    row += 1
    for item in VERSION_ROWS:
        ws.cell(row=row, column=1, value=item["component"])
        ws.cell(row=row, column=2, value=item["platform"])
        ws.cell(row=row, column=3, value=item["app_version"])
        ws.cell(row=row, column=4, value=item["build_artifact"])
        ws.cell(row=row, column=5, value=item["uat_url"])
        ws.cell(row=row, column=6, value=item["prod_url"])
        ws.cell(row=row, column=7, value=item["build_date"])
        ws.cell(row=row, column=8, value=item["remark"])
        for col in range(1, 9):
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="二、UAT 環境網址").font = Font(bold=True, size=12)
    row += 1
    for k, v in UAT_ENV.items():
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=v)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="三、本輪驗收範圍").font = Font(bold=True, size=12)
    row += 1
    for i, text in enumerate(SCOPE_ITEMS, start=1):
        ws.cell(row=row, column=1, value=f"{i}.")
        ws.cell(row=row, column=2, value=text)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="四、驗收結論").font = Font(bold=True, size=12)
    row += 1
    conclusions = [
        ("整體結果", "□ 通過　□ 有條件通過　□ 不通過"),
        ("CRM / 後台", "□ P　□ F　□ NA"),
        ("API", "□ P　□ F　□ NA"),
        ("Android 乘客", "□ P　□ F　□ NA"),
        ("Android 司機", "□ P　□ F　□ NA"),
        ("iOS 乘客", "□ P　□ F　□ NA"),
        ("iOS 司機", "□ P　□ F　□ NA"),
        ("阻塞問題 (P0)", ""),
        ("一般問題 (P1/P2)", ""),
        ("留待下版項目", ""),
    ]
    for col, h in enumerate(SIGNOFF_HEADERS, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        _header_style(cell, fg="4472C4")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    row += 1
    for a, b in conclusions:
        ws.cell(row=row, column=1, value=a)
        ws.cell(row=row, column=2, value=b)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="五、簽署").font = Font(bold=True, size=12)
    row += 2
    sign_rows = [
        ("測試負責人", "姓名：________________　日期：________________　簽名：________________"),
        ("開發／交付方", "姓名：________________　日期：________________　簽名：________________"),
        ("業務／客戶代表", "姓名：________________　日期：________________　簽名：________________"),
    ]
    for title, line in sign_rows:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True)
        ws.cell(row=row, column=2, value=line)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        row += 2

    widths = [14, 16, 10, 28, 22, 22, 12, 36]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: 對照驗收清單（可連結使用）
    ws2 = wb.create_sheet("版本對照")
    ws2.append(["說明", "詳細功能驗收請使用同目錄 TaxiMA-UAT-驗收清單.xlsx"])
    ws2.append(["產生簽收書日期", TODAY])

    wb.save(path)


def write_word(path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "PingFang TC"
    normal.font.size = Pt(11)

    title = doc.add_heading("Taxi MA（搭的）用戶驗收測試（UAT）簽收書", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(f"文件日期：{TODAY}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "本文件確認下列 CRM、後端 API、Android／iOS 乘客及司機應用之交付版本與 UAT 驗收結果。"
        "詳細測試步驟請參考《TaxiMA-UAT-驗收清單》。"
    )

    doc.add_heading("一、基本資料", level=1)
    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    rows_data = [
        ("專案名稱", "Taxi MA / 搭的"),
        ("驗收環境", "UAT（https://taxima.hkcodai.space）及 PROD 網址對照"),
        ("Monorepo 版本", "0.1.0-uat"),
        ("文件版本", "1.0"),
    ]
    for i, (k, v) in enumerate(rows_data):
        meta.rows[i].cells[0].text = k
        meta.rows[i].cells[1].text = v

    doc.add_heading("二、交付版本清單", level=1)
    table = doc.add_table(rows=1, cols=len(VERSION_HEADERS))
    table.style = "Table Grid"
    for i, h in enumerate(VERSION_HEADERS):
        table.rows[0].cells[i].text = h
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for item in VERSION_ROWS:
        row = table.add_row().cells
        row[0].text = item["component"]
        row[1].text = item["platform"]
        row[2].text = item["app_version"]
        row[3].text = item["build_artifact"]
        row[4].text = item["uat_url"]
        row[5].text = item["prod_url"]
        row[6].text = item["build_date"]
        row[7].text = item["remark"]

    doc.add_heading("三、UAT 環境網址", level=1)
    env_table = doc.add_table(rows=1, cols=2)
    env_table.style = "Table Grid"
    env_table.rows[0].cells[0].text = "項目"
    env_table.rows[0].cells[1].text = "URL"
    for k, v in UAT_ENV.items():
        r = env_table.add_row().cells
        r[0].text = k
        r[1].text = v

    doc.add_heading("四、本輪驗收範圍", level=1)
    for text in SCOPE_ITEMS:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("五、驗收結論", level=1)
    doc.add_paragraph("整體結果：□ 通過　□ 有條件通過　□ 不通過")
    conc = doc.add_table(rows=1, cols=3)
    conc.style = "Table Grid"
    conc.rows[0].cells[0].text = "模組"
    conc.rows[0].cells[1].text = "結果 (P/F/NA)"
    conc.rows[0].cells[2].text = "備註"
    for mod in [
        "CRM / 後台",
        "API",
        "Android 乘客 App",
        "Android 司機 App",
        "iOS 乘客 App",
        "iOS 司機 App",
    ]:
        r = conc.add_row().cells
        r[0].text = mod
        r[1].text = ""
        r[2].text = ""

    doc.add_paragraph()
    doc.add_paragraph("阻塞問題 (P0)：")
    doc.add_paragraph("一般問題 (P1/P2)：")
    doc.add_paragraph("留待下版項目：")

    doc.add_page_break()
    doc.add_heading("六、簽署", level=1)
    doc.add_paragraph(
        "各方代表確認已於 UAT 環境完成上述範圍之測試，並同意本文件所載版本為驗收標的。"
    )
    for role in ["測試負責人", "開發／交付方", "業務／客戶代表"]:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run(f"{role}\n").bold = True
        doc.add_paragraph("姓名：_______________________________")
        doc.add_paragraph("日期：_______________________________")
        doc.add_paragraph("簽名：_______________________________")

    doc.save(path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx = OUT_DIR / "TaxiMA-UAT-簽收書.xlsx"
    docx = OUT_DIR / "TaxiMA-UAT-簽收書.docx"
    write_excel(xlsx)
    write_word(docx)
    print(xlsx)
    print(docx)


if __name__ == "__main__":
    main()
