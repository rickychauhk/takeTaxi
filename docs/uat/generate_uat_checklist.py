#!/usr/bin/env python3
"""Generate UAT checklist Excel + Word for Taxi MA."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent
TODAY = date.today().isoformat()

HEADERS = ["類別", "功能", "驗收標準", "結果(P/F/NA)", "測試人", "日期", "備註"]

ROWS: list[tuple[str, str, str]] = [
    # 0. 準備
    ("0. UAT 前準備", "API 健康", "GET /api/v1/health 回 ok:true，env:uat"),
    ("0. UAT 前準備", "CRM 可登入", "管理員登入／登出正常"),
    ("0. UAT 前準備", "測試帳號", "≥2 乘客、≥3 司機（含已審批、1 未審批、1 可出貨）"),
    ("0. UAT 前準備", "App 版本", "記錄乘客／司機 build 與 API 部署時間"),
    ("0. UAT 前準備", "推播", "OneSignal 已允許通知（前景／背景各測一次）"),
    ("0. UAT 前準備", "定位", "乘客／司機已允許定位（司機含背景若需測）"),
    ("0. UAT 前準備", "測試資料", "準備過海／陸路、即時／預約、積分足夠乘客各一組"),
    # 1. CRM — 營運
    ("1. CRM", "儀表板", "首頁載入無錯誤；關鍵數字合理"),
    ("1. CRM", "即時柯打", "狀態流轉即時；可開詳情 drawer"),
    ("1. CRM", "乘客落柯打紀錄", "篩選、分頁、詳情與 App 一致"),
    ("1. CRM", "司機接單紀錄", "接單時間、司機、距離、狀態正確"),
    ("1. CRM", "行程對話紀錄", "訂單詳情可查看 chat transcript"),
    ("1. CRM", "評分", "星級、選填理由、關聯訂單可見"),
    ("1. CRM", "司機列表", "搜尋、分頁、進入詳情"),
    ("1. CRM", "司機審批", "批准／拒絕後 App 權限跟隨"),
    ("1. CRM", "司機文件", "行車證／司機證可檢視"),
    ("1. CRM", "乘客列表", "會員 ID、電話、積分、編輯"),
    ("1. CRM", "意見反饋", "乘客提交後列表可見、可讀內容"),
    # 1. CRM — 設定與內容
    ("1. CRM", "叫車選項", "分類／選項 CRUD；乘客「更多選項」同步"),
    ("1. CRM", "乘客附近司機範圍", "改公里數後乘客地圖／等候接單半徑跟隨"),
    ("1. CRM", "乘客行程積分", "兌換比例、開關；完成行程加點邏輯"),
    ("1. CRM", "司機積分套餐", "套餐價格與司機 App 購點一致"),
    ("1. CRM", "積分購買審批", "批准／拒絕後司機點數更新"),
    ("1. CRM", "司機繁忙時段", "設定時段後大廳／接單點數加成反映"),
    ("1. CRM", "訊息中心", "發推播；乘客 App「訊息」列表與詳情"),
    ("1. CRM", "聯絡我們（網頁）", "HTML 儲存後乘客 App 顯示正確"),
    ("1. CRM", "App 使用條款", "儲存後乘客／司機 WebView 正確"),
    ("1. CRM", "公司價告示", "儲存後乘客確認叫車「公司價」文案正確"),
    ("1. CRM", "最新消息連結", "側欄或設定內連結可開（如有配置）"),
    ("1. CRM", "乘客快捷訊息", "維護後行程 chat 可選快捷句"),
    ("1. CRM", "優惠碼", "建立／停用；司機註冊兌換成功或拒絕理由"),
    ("1. CRM", "廣告管理", "上架廣告在 App 指定位置顯示（如有）"),
    ("1. CRM", "APK 發佈", "上傳版本；下載連結或版本資訊正確（如使用）"),
    ("1. CRM", "Gemini 設定", "儲存不報錯；關聯 AI 功能可用（如有開）"),
    ("1. CRM", "系統日誌", "敏感操作（審批、設定）有稽核記錄"),
    # 2. 乘客 App — 登入與首頁
    ("2. 乘客 App", "電話登入", "8 位本地號；新舊會員；錯號提示"),
    ("2. 乘客 App", "記住電話", "勾選後下次預填"),
    ("2. 乘客 App", "側欄", "會員 ID、電話 8 位無 +852；各入口可開"),
    ("2. 乘客 App", "GPS 定位", "首次權限；地圖中心合理"),
    ("2. 乘客 App", "搜尋地點", "起／終點關鍵字搜尋選點"),
    ("2. 乘客 App", "地圖揀點", "拖 pin／地圖選點作起終點"),
    ("2. 乘客 App", "附近司機（叫車前）", "地圖有司機 marker；範圍跟 CRM"),
    ("2. 乘客 App", "路線預覽", "起終點連線、距離時間、隧道費估算"),
    ("2. 乘客 App", "過海隧道", "可選偏好隧道；路線／費用跟隨"),
    # 2. 乘客 App — 確認叫車
    ("2. 乘客 App", "確認叫車—車資", "市區的士估算、隧道附加說明"),
    ("2. 乘客 App", "確認叫車—小費", "開關與金額帶入訂單"),
    ("2. 乘客 App", "確認叫車—六人座", "選項影響叫車參數"),
    ("2. 乘客 App", "確認叫車—積分抵扣", "餘額、上限、折抵金額正確"),
    ("2. 乘客 App", "確認叫車—付款方式", "多選合法；提交帶入"),
    ("2. 乘客 App", "確認叫車—聯絡方式", "致電／SMS 選項帶入"),
    ("2. 乘客 App", "確認叫車—公司價", "勾選彈窗與 CRM 告示一致"),
    ("2. 乘客 App", "更多選項", "CRM 附加選項、備註、即時／預約時間"),
    ("2. 乘客 App", "預約叫車", "選未來時間；確認頁顯示預約時間"),
    # 2. 乘客 App — 行程中
    ("2. 乘客 App", "等候接單", "附近司機地圖；不可改隧道；可取消"),
    ("2. 乘客 App", "司機接單", "司機資訊、車牌、ETA、路線"),
    ("2. 乘客 App", "司機前往", "地圖跟隨；可致電司機"),
    ("2. 乘客 App", "行程中", "狀態 en_route／in_trip；路線可展開／收起"),
    ("2. 乘客 App", "行程 chat", "互傳文字；未讀；快捷訊息（如有）"),
    ("2. 乘客 App", "乘客取消（等候）", "取消成功；不殘留 ghost 訂單"),
    ("2. 乘客 App", "司機取消", "提示重叫；chat 關閉"),
    ("2. 乘客 App", "完成評分", "星級＋選填理由；可略過"),
    ("2. 乘客 App", "完成頁", "車資、小費、積分使用／獲得、結餘"),
    # 2. 乘客 App — 其他
    ("2. 乘客 App", "叫車紀錄", "列表狀態、詳情、積分／小費欄位"),
    ("2. 乘客 App", "訊息", "推播列表、詳情、已讀"),
    ("2. 乘客 App", "失物協助—入口", "側欄有；設定內無"),
    ("2. 乘客 App", "失物協助—提交", "表單、相簿權限、多圖上傳"),
    ("2. 乘客 App", "失物協助—下架", "約 20 分鐘推播＋訊息列表提示"),
    ("2. 乘客 App", "意見反映", "提交成功；CRM 可見"),
    ("2. 乘客 App", "聯絡我們", "顯示 CRM HTML"),
    ("2. 乘客 App", "設定", "語言、條款、版本；無訊息／失物入口"),
    ("2. 乘客 App", "推播點擊", "點通知可進 App 對應畫面（訊息／訂單）"),
    # 3. 司機 App — 帳號
    ("3. 司機 App", "登入", "已註冊電話登入"),
    ("3. 司機 App", "註冊", "電話、車牌、上傳行車證／司機證"),
    ("3. 司機 App", "註冊優惠碼", "有效碼加點；無效碼有理由"),
    ("3. 司機 App", "待審批", "未審批不可上線接單／出貨"),
    ("3. 司機 App", "補交證件", "上傳失敗後主畫面可重試"),
    # 3. 司機 App — 大廳與接單
    ("3. 司機 App", "上線／下線", "上線收定位；下線不收單"),
    ("3. 司機 App", "接單半徑", "超出 CRM／設定半徑不接"),
    ("3. 司機 App", "大廳列表", "距離、起終點、附加、預約標示"),
    ("3. 司機 App", "大廳排序", "最新／最近預約排序切換"),
    ("3. 司機 App", "繁忙加成", "繁忙時段訂單點數加成顯示"),
    ("3. 司機 App", "接即時單", "接單→前往→開始→完成"),
    ("3. 司機 App", "接預約單", "接受後入預約 Tab；提示時間"),
    ("3. 司機 App", "預約 Tab", "列表、放回大廳、開始行程時間限制"),
    ("3. 司機 App", "司機取消", "取消後乘客端同步"),
    ("3. 司機 App", "行程 chat／致電", "與乘客對話、撥號"),
    ("3. 司機 App", "失物走馬燈（有）", "頂部捲動；多則 ◆ 分隔"),
    ("3. 司機 App", "失物走馬燈（無）", "無失物時整條不顯示"),
    ("3. 司機 App", "失物 WebView", "描述、圖片、聯絡方式"),
    # 3. 司機 App — 出貨與點數
    ("3. 司機 App", "出貨—選點", "搜尋／地圖選起終點"),
    ("3. 司機 App", "出貨—選項", "隧道、附加、預約與乘客叫車一致邏輯"),
    ("3. 司機 App", "出貨—發佈", "扣點／出現在他人大廳"),
    ("3. 司機 App", "今日完成", "列表與收入統計"),
    ("3. 司機 App", "購點套餐", "選套餐、上傳入數憑證"),
    ("3. 司機 App", "設定", "語言、條款、版本、登出"),
    # 4. 叫車·接單專項（乘客落單 × 司機接單）
    ("4. 叫車·接單", "前置：雙司機上線", "≥2 司機上線、定位在上車點附近；CRM 即時柯打可見 matching"),
    ("4. 叫車·接單", "P-01 陸路即時（基本）", "乘客：短程同區起終點→確認叫車→等候；司機：大廳見單→接單→完成"),
    ("4. 叫車·接單", "P-02 過海即時", "乘客：選過海隧道偏好；路線過海；司機：大廳顯示隧道／附加；接單至完成"),
    ("4. 叫車·接單", "P-03 陸路隧道", "乘客：需經獅子山／大老山等陸隧路段；估算含隧道附加；司機可接完成"),
    ("4. 叫車·接單", "P-04 無過海隧道", "乘客：純陸路（如大埔道）；隧道費為 0；司機接單正常"),
    ("4. 叫車·接單", "P-05 小費", "乘客：確認頁加小費；完成頁／叫車紀錄顯示小費；CRM 詳情一致"),
    ("4. 叫車·接單", "P-06 六人座", "乘客：勾六人座叫車；司機大廳或詳情可辨識；完成無報錯"),
    ("4. 叫車·接單", "P-07 積分抵扣", "乘客：用部分積分；完成扣點、顯示折抵；CRM 乘客積分更新"),
    ("4. 叫車·接單", "P-08 付款方式", "乘客：多選付款方式叫車；司機／CRM 可見選項摘要"),
    ("4. 叫車·接單", "P-09 公司價", "乘客：勾公司價並確認告示；司機接單→完成"),
    ("4. 叫車·接單", "P-10 備註", "乘客：更多選項填備註；司機大廳 extra 行或詳情可見"),
    ("4. 叫車·接單", "P-11 CRM 附加選項", "乘客：多選 CRM 叫車選項；司機端顯示一致"),
    ("4. 叫車·接單", "P-12 聯絡偏好", "乘客：只選致電或 SMS；行程中致電按鈕可用"),
    ("4. 叫車·接單", "P-13 預約 30–60 分鐘", "乘客：選近 future 時間；狀態 matching；司機接「預約」入 Tab"),
    ("4. 叫車·接單", "P-14 預約 隔日", "乘客：選次日時段；司機大廳有預約時間標示；排序「最近預約」"),
    ("4. 叫車·接單", "P-15 等候 UI", "乘客 matching：地圖見附近司機；隧道 picker 唯讀；可取消"),
    ("4. 叫車·接單", "P-16 接單後乘客", "司機接單後：乘客見車牌／司機名；ETA／路線；司機位置更新"),
    ("4. 叫車·接單", "P-17 狀態 accepted", "司機接單未出發：乘客 accepted；司機可「前往上車點」"),
    ("4. 叫車·接單", "P-18 狀態 en_route", "司機 en_route：乘客文案／地圖跟隨；CRM 即時柯打更新"),
    ("4. 叫車·接單", "P-19 狀態 in_trip", "司機開始行程 in_trip：乘客顯示行程中；終點路線為下車點"),
    ("4. 叫車·接單", "P-20 完成評分", "司機 completed：乘客評分頁；可提交或略過；再回地圖可叫新單"),
    ("4. 叫車·接單", "P-21 重開 App", "行程進行中 kill 乘客 App 再開：仍進 ActiveOrder；狀態與 API 一致"),
    ("4. 叫車·接單", "P-22 進行中不可重叫", "有 active 行程時不可再落第二張即時單（或依產品規則提示）"),
    ("4. 叫車·接單", "D-01 接單半徑內", "司機設定 X km；上車點在內：大廳列表有單、可接"),
    ("4. 叫車·接單", "D-02 接單半徑外", "司機在上車點 X km 外：列表無該單或接單被阻擋"),
    ("4. 叫車·接單", "D-03 大廳手動接單", "司機：點列表→確認→進 ActiveTrip；乘客同步接單"),
    ("4. 叫車·接單", "D-04 搶單（雙司機）", "兩司機同時接同一 matching：只得一人成功；另一人提示失敗／單消失"),
    ("4. 叫車·接單", "D-05 背景新單提示", "司機 App 背景／另一 App 前景：新單 push 或全局 alert；可接或拒"),
    ("4. 叫車·接單", "D-06 繁忙時段接單", "CRM 繁忙時段內落單：司機大廳點數／加成與設定一致；接單扣點正確"),
    ("4. 叫車·接單", "D-07 預約 reserve", "司機接預約：提示已接受；訂單入預約 Tab；大廳即時單列表不含該單"),
    ("4. 叫車·接單", "D-08 預約過早開始", "距預約時間尚早：司機「開始行程」被擋；提示時間"),
    ("4. 叫車·接單", "D-09 預約準時開始", "到可開始窗口：司機 in_trip；乘客 in_trip"),
    ("4. 叫車·接單", "D-10 預約放回大廳", "司機將預約放回：乘客端／CRM 狀態回 matching；其他司機可見"),
    ("4. 叫車·接單", "D-11 司機 misaccept", "司機接單後取消（misaccept）：乘客可重叫；紀錄「司機曾取消接單」"),
    ("4. 叫車·接單", "D-12 司機取消（已接）", "accepted／en_route 司機取消：乘客提示；chat 不可用；CRM 紀錄取消方"),
    ("4. 叫車·接單", "D-13 乘客等候取消", "matching 乘客取消：司機大廳移除；CRM cancelled passenger"),
    ("4. 叫車·接單", "D-14 下線不接單", "司機下線：不收新單；進行中行程不受影響"),
    ("4. 叫車·接單", "D-15 未審批不可接", "未審批司機：大廳無單或接單被拒"),
    ("4. 叫車·接單", "D-16 已有 active 單", "司機已有一張進行中：不可再接第二張（或依規則提示）"),
    ("4. 叫車·接單", "S-01 司機出貨（無懸賞）", "司機 A 出貨：司機 B 大廳見乘客風格／出貨標記；B 接單→完成"),
    ("4. 叫車·接單", "S-02 出貨含懸賞", "司機 A 出貨設懸賞點數：A 扣點；B 接單完成後 B 得懸賞（推播／點數）"),
    ("4. 叫車·接單", "S-03 出貨 matching 取消", "出貨單無人接前 A 或乘客端取消：懸賞退回 A"),
    ("4. 叫車·接單", "S-04 出貨預約", "司機出貨選預約時間：B 接預約流程同 P-13"),
    ("4. 叫車·接單", "C-01 行程 chat 後接單", "matching 時乘客 chat（若允許）或接單後 chat：雙向收訊；CRM transcript"),
    ("4. 叫車·接單", "C-02 chat 放回大廳", "司機／乘客 chat 觸發 release_to_hall（如有）：單回大廳、狀態正確"),
    ("4. 叫車·接單", "CRM-01 即時柯打", "落單至完成：CRM Live 狀態 pending→matching→accepted→…→completed"),
    ("4. 叫車·接單", "CRM-02 雙向紀錄", "同一 orderId：乘客落柯打、司機接單紀錄時間／距離一致"),
    ("4. 叫車·接單", "CRM-03 取消軌跡", "各取消 case：CRM 詳情 status events 與 actor 正確"),
    # 5. 端到端（跨模組 smoke）
    ("5. 端到端", "A 即時叫車", "叫車→接單→完成→評分；CRM／紀錄一致"),
    ("5. 端到端", "B 預約叫車", "乘客預約→司機接→Tab→到時開始→完成"),
    ("5. 端到端", "C 乘客取消", "等候取消；司機大廳移除"),
    ("5. 端到端", "D 司機取消", "乘客重叫；chat 關閉"),
    ("5. 端到端", "E 司機出貨", "A 出貨→B 接單→完成"),
    ("5. 端到端", "F 失物協助", "提交→走馬燈→詳情→下架推播"),
    ("5. 端到端", "G CRM 推播", "全員或指定；訊息 inbox"),
    ("5. 端到端", "H 積分", "完成加點；下單抵扣；CRM 餘額一致"),
    ("5. 端到端", "I 過海隧道", "選隧道→路線／帳單隧道費合理"),
    ("5. 端到端", "J 雙 chat", "乘客司機互發；CRM transcript 完整"),
    # 6. 非功能
    ("6. 非功能", "繁中／英文", "乘客／司機設定切換主要文案"),
    ("6. 非功能", "弱網", "斷網提示；恢復可重試叫車／接單"),
    ("6. 非功能", "Android SSL", "UAT API 無 certificate handshake 錯誤"),
    ("6. 非功能", "iOS／Android", "各測至少一台真機（乘客＋司機）"),
    ("6. 非功能", "背景定位", "司機上線後 App 背景仍更新位置（抽測）"),
]

BOOKING_ACCEPT_PREFIX = "4. 叫車·接單"


def _fill_checklist_sheet(ws, sheet_title: str, rows: list[tuple[str, str, str]]) -> None:
    title = f"{sheet_title}（產生日期：{TODAY}）"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal="center")

    ws.append([])
    header_row = 3
    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C00000")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, (cat, feat, crit) in enumerate(rows, start=header_row + 1):
        ws.cell(row=i, column=1, value=cat)
        ws.cell(row=i, column=2, value=feat)
        ws.cell(row=i, column=3, value=crit)
        for col in range(1, 8):
            ws.cell(row=i, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 48
    for col in "DEFG":
        ws.column_dimensions[col].width = 14


def write_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "UAT驗收清單"
    _fill_checklist_sheet(ws, "Taxi MA UAT 驗收清單", ROWS)

    booking_rows = [r for r in ROWS if r[0] == BOOKING_ACCEPT_PREFIX]
    ws_b = wb.create_sheet("叫車接單專項")
    _fill_checklist_sheet(ws_b, "Taxi MA 叫車·接單專項", booking_rows)

    meta = wb.create_sheet("測試記錄")
    meta["A1"] = "欄位"
    meta["B1"] = "填寫"
    meta["A1"].font = Font(bold=True)
    meta["B1"].font = Font(bold=True)
    fields = [
        ("測試日期", ""),
        ("CRM 版本／部署", ""),
        ("API health 時間戳", ""),
        ("乘客 App build", ""),
        ("司機 App build", ""),
        ("UAT 網址", "https://taxima.hkcodai.space"),
        ("阻塞問題 P0", ""),
        ("一般問題 P1/P2", ""),
    ]
    for r, (k, v) in enumerate(fields, start=2):
        meta.cell(row=r, column=1, value=k)
        meta.cell(row=r, column=2, value=v)
    meta.column_dimensions["A"].width = 22
    meta.column_dimensions["B"].width = 60

    wb.save(path)


def write_word(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "PingFang TC"
    style.font.size = Pt(11)

    t = doc.add_heading("Taxi MA UAT 驗收清單", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"產生日期：{TODAY}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "環境建議：UAT — https://taxima.hkcodai.space（API/CRM）；請使用最新乘客／司機 App 建置。"
    )

    from itertools import groupby

    for cat, group in groupby(ROWS, key=lambda r: r[0]):
        doc.add_heading(cat, level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "功能"
        hdr[1].text = "驗收標準"
        hdr[2].text = "結果(P/F/NA)"
        hdr[3].text = "備註"
        for cell in hdr:
            for par in cell.paragraphs:
                for run in par.runs:
                    run.bold = True
        for _cat, feat, crit in group:
            row = table.add_row().cells
            row[0].text = feat
            row[1].text = crit
            row[2].text = ""
            row[3].text = ""

    doc.add_page_break()
    doc.add_heading("測試記錄", level=1)
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.style = "Table Grid"
    meta_table.rows[0].cells[0].text = "欄位"
    meta_table.rows[0].cells[1].text = "填寫"
    for k, v in [
        ("測試日期", ""),
        ("CRM 版本／部署", ""),
        ("API health 時間戳", ""),
        ("乘客 App build", ""),
        ("司機 App build", ""),
        ("UAT 網址", "https://taxima.hkcodai.space"),
        ("阻塞問題 P0", ""),
        ("一般問題 P1/P2", ""),
    ]:
        r = meta_table.add_row().cells
        r[0].text = k
        r[1].text = v

    doc.save(path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx = OUT_DIR / "TaxiMA-UAT-驗收清單.xlsx"
    docx = OUT_DIR / "TaxiMA-UAT-驗收清單.docx"
    write_excel(xlsx)
    write_word(docx)
    print(xlsx)
    print(docx)


if __name__ == "__main__":
    main()
